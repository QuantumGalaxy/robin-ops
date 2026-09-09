"""DoltHub research history. Kept separate from the execution IV archive.

The vendor's IV tenor and observation-date conventions are not yet verified.
Importing data must never silently enable its use by the entry screener.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import sqlite3
import ssl
import urllib.parse
import urllib.request
from contextlib import closing
from datetime import UTC, date, datetime
from pathlib import Path

import certifi

from .reference_feed import calendar, last_completed

SOURCE = "DoltHub post-no-preference/options volatility_history.iv_current"
DATABASE = "iv-research.sqlite3"


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("Excel import requires optionsagent[excel]") from None
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            populated = [s for s in book if s.max_row > 1]
            if len(populated) != 1:
                raise ValueError("Expected one populated worksheet; export sheets separately")
            records = iter(populated[0].iter_rows(values_only=True))
            headers = [str(x).strip() for x in next(records)]
            values = list(records)
        finally:
            book.close()
    elif path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            records = iter(csv.reader(f))
            headers = next(records, [])
            values = list(records)
    else:
        raise ValueError("Use a CSV or XLSX file")
    aliases = {"act_symbol": "symbol", "iv_current": "iv"}
    headers = [aliases.get(x.strip(), x.strip()) for x in headers]
    if len(set(headers)) != len(headers) or not {"symbol", "date", "iv"} <= set(headers):
        raise ValueError("Expected unique symbol,date,iv columns (DoltHub names also accepted)")
    return [
        dict(zip(headers, r, strict=True))
        for r in values
        if any(x is not None and x != "" for x in r)
    ]


def connect(path):
    db = sqlite3.connect(path)
    db.executescript("""
        CREATE TABLE IF NOT EXISTS observations(
            symbol TEXT, day TEXT, iv REAL, source TEXT, file_hash TEXT,
            PRIMARY KEY(symbol, day));
        CREATE TABLE IF NOT EXISTS imports(
            file_hash TEXT PRIMARY KEY, imported_at TEXT, filename TEXT, report TEXT);
        CREATE TABLE IF NOT EXISTS rejected(
            file_hash TEXT, row_number INTEGER, reason TEXT, raw TEXT,
            PRIMARY KEY(file_hash, row_number));
    """)
    return db


def import_history(path, state_dir, symbols, now=None):
    """Atomic, repeatable import with rejected rows preserved for review."""
    now = now or datetime.now(UTC)
    path, state_dir = Path(path), Path(state_dir)
    records = read_rows(path)
    if not records:
        raise ValueError("No observations found")
    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    completed = last_completed(now)
    cal = calendar()
    state_dir.mkdir(parents=True, exist_ok=True)
    with closing(connect(state_dir / DATABASE)) as db, db:
        previous = db.execute(
            "SELECT report FROM imports WHERE file_hash=?", (file_hash,)
        ).fetchone()
        if previous:
            return {**json.loads(previous[0]), "already_imported": True}
        accepted = rejected = duplicates = 0
        for number, row in enumerate(records, 2):
            try:
                symbol = str(row["symbol"]).strip().upper()
                if symbol not in symbols:
                    raise ValueError("symbol outside configured watchlist")
                raw_day = row["date"]
                if isinstance(raw_day, datetime):
                    if raw_day.time().isoformat() != "00:00:00":
                        raise ValueError("date contains a time; observation convention unknown")
                    day = raw_day.date()
                else:
                    day = raw_day if isinstance(raw_day, date) else date.fromisoformat(str(raw_day))
                value = float(row["iv"])
                if isinstance(row["iv"], bool) or not math.isfinite(value) or value <= 0:
                    raise ValueError("IV must be finite and positive")
                if day > completed:
                    raise ValueError("future or incomplete session")
                if not cal.is_session(day.isoformat()):
                    raise ValueError("non-trading date; not relabeled")
                existing = db.execute(
                    "SELECT iv,source FROM observations WHERE symbol=? AND day=?",
                    (symbol, day.isoformat()),
                ).fetchone()
                if existing:
                    if existing != (value, SOURCE):
                        raise ValueError("conflicting existing observation; not overwritten")
                    duplicates += 1
                    continue
                db.execute(
                    "INSERT INTO observations VALUES (?,?,?,?,?)",
                    (symbol, day.isoformat(), value, SOURCE, file_hash),
                )
                accepted += 1
            except (ValueError, TypeError, OverflowError) as exc:
                db.execute(
                    "INSERT INTO rejected VALUES (?,?,?,?)",
                    (file_hash, number, str(exc), json.dumps(row, default=str)),
                )
                rejected += 1
        report = {
            "source": SOURCE,
            "rows": len(records),
            "accepted": accepted,
            "rejected": rejected,
            "duplicates": duplicates,
            "sha256": file_hash,
            "execution_enabled": False,
            "methodology": "unverified",
        }
        db.execute(
            "INSERT INTO imports VALUES (?,?,?,?)",
            (file_hash, now.isoformat(), path.name, json.dumps(report)),
        )
    return report


def _rank_details(rows, completed):
    result = {
        "rank": None,
        "days": len(rows),
        "latest": rows[-1][0] if rows else None,
        "source": SOURCE,
        "reason": "",
    }
    if len(rows) < 200:
        result["reason"] = "fewer than 200 observations in the last 252 sessions"
    elif rows[-1][0] != completed.isoformat():
        result["reason"] = "latest completed-session IV missing"
    elif any(
        source != SOURCE or not math.isfinite(value) or value <= 0 for _, value, source in rows
    ):
        result["reason"] = "invalid values or inconsistent IV source"
    else:
        values = [r[1] for r in rows]
        lo, hi = min(values), max(values)
        if lo == hi:
            result["reason"] = "flat history cannot establish IV rank"
        else:
            result["rank"] = (values[-1] - lo) / (hi - lo)
            result["reason"] = "experimental completed-day IV rank available"
    return result


def paper_iv_rank(state_dir, symbol, now=None):
    """Same-source completed-day rank; no current-contract IV substitution or stale fallback."""
    completed = last_completed(now or datetime.now(UTC))
    path = Path(state_dir) / DATABASE
    if not path.exists():
        return _rank_details([], completed)
    start = calendar().session_offset(completed.isoformat(), -251).date().isoformat()
    with closing(sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)) as db:
        rows = db.execute(
            "SELECT day,iv,source FROM observations WHERE symbol=? "
            "AND day>=? AND day<=? ORDER BY day",
            (symbol, start, completed.isoformat()),
        ).fetchall()
    return _rank_details(rows, completed)


def fetch_latest(symbols, completed):
    """Small public read query; never accept timed-out or row-limited results."""
    if not symbols or any(not s.isascii() or not s.isalnum() for s in symbols):
        raise ValueError("Expected simple US ticker symbols")
    quoted = ",".join("'" + s + "'" for s in symbols)
    query = (
        "SELECT date, act_symbol, iv_current FROM volatility_history "
        f"WHERE date = '{completed.isoformat()}' AND act_symbol IN ({quoted})"
    )
    url = "https://www.dolthub.com/api/v1alpha1/post-no-preference/options/master?"
    request = urllib.request.Request(
        url + urllib.parse.urlencode({"q": query}), headers={"Accept": "application/json"}
    )
    with urllib.request.urlopen(
        request, timeout=20, context=ssl.create_default_context(cafile=certifi.where())
    ) as response:
        raw = response.read(1_048_577)
    if len(raw) > 1_048_576:
        raise ValueError("IV response too large")
    result = json.loads(raw)
    if not isinstance(result, dict) or result.get("query_execution_status") != "Success":
        raise ValueError("IV query incomplete or failed")
    rows = result.get("rows")
    if not isinstance(rows, list) or len(rows) > len(symbols):
        raise ValueError("Unexpected IV row count")
    seen = set()
    for row in rows:
        if (
            not isinstance(row, dict)
            or row.get("date") != completed.isoformat()
            or row.get("act_symbol") not in symbols
            or row["act_symbol"] in seen
        ):
            raise ValueError("Unexpected IV symbol, date, or duplicate")
        seen.add(row["act_symbol"])
    return result


def refresh_latest(state_dir, symbols, now=None, fetcher=fetch_latest):
    now = now or datetime.now(UTC)
    completed = last_completed(now)
    status = history_status(state_dir, symbols, now)
    fresh = {s["symbol"] for s in status["symbols"] if s["fresh"]}
    missing = [s for s in symbols if s not in fresh]
    if not missing:
        return {"message": "Completed-day IV already cached", "required_through": str(completed)}
    result = fetcher(missing, completed)
    folder = Path(state_dir) / "iv-downloads"
    folder.mkdir(parents=True, exist_ok=True)
    encoded = json.dumps(result, sort_keys=True).encode()
    stem = f"{completed}-{hashlib.sha256(encoded).hexdigest()[:16]}"
    (folder / (stem + ".json")).write_bytes(encoded)
    path = folder / (stem + ".csv")
    with path.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["symbol", "date", "iv"])
        for row in result["rows"]:
            writer.writerow([row["act_symbol"], row["date"], row.get("iv_current")])
    if not result["rows"]:
        return {
            "status": "pending",
            "required_through": str(completed),
            "missing_symbols": missing,
            "message": "DoltHub has not published the required completed-session IV",
        }
    imported = import_history(path, state_dir, symbols, now)
    remaining = [
        r["symbol"] for r in history_status(state_dir, symbols, now)["symbols"] if not r["fresh"]
    ]
    return {
        **imported,
        "status": "pending" if remaining else "complete",
        "required_through": str(completed),
        "missing_symbols": remaining,
    }


def history_status(state_dir, symbols, now=None, experimental=False):
    path = Path(state_dir) / DATABASE
    if not path.exists():
        return {"status": "Not imported", "symbols": []}
    completed = last_completed(now or datetime.now(UTC))
    cal = calendar()
    end = cal.date_to_session(completed.isoformat())
    start = cal.session_offset(end, -251)
    expected = {d.date().isoformat() for d in cal.sessions_in_range(start, end)}
    summaries = []
    with closing(sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)) as db:
        for symbol in symbols:
            rows = db.execute(
                "SELECT day,iv,source FROM observations WHERE symbol=? "
                "AND day>=? AND day<=? ORDER BY day",
                (symbol, min(expected), completed.isoformat()),
            ).fetchall()
            dates = {r[0] for r in rows}
            summaries.append(
                {
                    "symbol": symbol,
                    "days": len(rows),
                    "latest": rows[-1][0] if rows else None,
                    "missing_sessions": sorted(expected - dates),
                    "fresh": completed.isoformat() in dates,
                    "paper_rank": _rank_details(rows, completed)["rank"],
                    "rank_reason": _rank_details(rows, completed)["reason"],
                }
            )
        rejected = db.execute("SELECT COUNT(*) FROM rejected").fetchone()[0]
    return {
        "status": (
            "Paper IV filter ON — experimental, source methodology unverified"
            if experimental
            else "Research only — IV definition and date labels unverified"
        ),
        "paper_filter_enabled": experimental,
        "source": SOURCE,
        "required_through": completed.isoformat(),
        "execution_enabled": False,
        "rejected_rows": rejected,
        "symbols": summaries,
    }
